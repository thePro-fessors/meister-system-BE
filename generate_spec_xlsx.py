import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
wb.remove(wb.active)  # remove default sheet

# Color Palette (Professional Slate & Swiss Theme)
C_NAVY_HEADER = "1E293B"     # Slate 800
C_NAVY_TEXT = "FFFFFF"
C_SUBHEADER = "334155"       # Slate 700
C_SECTION_BG = "F1F5F9"      # Slate 100
C_ALT_ROW = "F8FAFC"         # Slate 50
C_BORDER = "CBD5E1"          # Slate 300
C_ACCENT_RED = "E30613"      # Swiss Red
C_GREEN_BG = "DCFCE7"        # Green 100
C_GREEN_TXT = "166534"
C_BLUE_BG = "DBEAFE"         # Blue 100
C_BLUE_TXT = "1E40AF"
C_AMBER_BG = "FEF3C7"        # Amber 100
C_AMBER_TXT = "92400E"
C_PURPLE_BG = "F3E8FF"       # Purple 100
C_PURPLE_TXT = "6B21A8"

thin_border = Border(
    left=Side(style='thin', color=C_BORDER),
    right=Side(style='thin', color=C_BORDER),
    top=Side(style='thin', color=C_BORDER),
    bottom=Side(style='thin', color=C_BORDER)
)

header_font = Font(name="Pretendard", size=10.5, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color=C_NAVY_HEADER, end_color=C_NAVY_HEADER, fill_type="solid")

sub_header_font = Font(name="Pretendard", size=10, bold=True, color="FFFFFF")
sub_header_fill = PatternFill(start_color=C_SUBHEADER, end_color=C_SUBHEADER, fill_type="solid")

title_font = Font(name="Pretendard", size=15, bold=True, color="0F172A")
subtitle_font = Font(name="Pretendard", size=9.5, color="64748B")

data_font = Font(name="Pretendard", size=9.5, color="0F172A")
data_font_bold = Font(name="Pretendard", size=9.5, bold=True, color="0F172A")
data_font_code = Font(name="Consolas", size=9.5, color="0F172A")
data_font_small = Font(name="Pretendard", size=8.5, color="475569")

def style_table(ws, min_col, max_col, start_row, end_row, zebra=True):
    for r in range(start_row, end_row + 1):
        is_even = (r % 2 == 0)
        for c in range(min_col, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = thin_border
            if zebra and is_even and not cell.fill.start_color.rgb:
                cell.fill = PatternFill(start_color=C_ALT_ROW, end_color=C_ALT_ROW, fill_type="solid")

def auto_width(ws, max_widths=None):
    if max_widths is None:
        max_widths = {}
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = str(cell.value or "")
            lines = val.split("\n")
            line_len = max(len(l.encode('utf-8')) for l in lines) if lines else 0
            if line_len > max_len:
                max_len = line_len
        # Convert byte len approx to col width
        approx_w = max(10, min(max_widths.get(col_letter, 50), int(max_len * 0.75) + 3))
        ws.column_dimensions[col_letter].width = approx_w

# ==============================================================================
# SHEET 1: 01_시스템_개요 (Overview)
# ==============================================================================
ws1 = wb.create_sheet(title="01_시스템_개요")
ws1.views.sheetView[0].showGridLines = True

ws1.append(["부산소프트웨어마이스터고 마이스터 역량인증제 - 시스템 전체 기능 & API 명세서"])
ws1.append(["생성일시: 2026-10-15 | 백엔드: FastAPI (Python 3.12) + MySQL 8.0 + Redis 7.0 | 작성: Antigravity"])
ws1.append([])

ws1["A1"].font = title_font
ws1["A2"].font = subtitle_font

overview_sections = [
    ("1. 기본 통신 및 환경 정보", [
        ("로컬 개발 API 베이스", "http://localhost/api 또는 http://127.0.0.1:8000/api", "Nginx(80) 리버스 프록시를 통해 백엔드(8000) 및 프론트엔드(3000) 통합 서빙"),
        ("Swagger / OpenAPI 문서", "http://localhost/docs 및 http://localhost/redoc", "개발 환경 전체 대화형 API 문서 제공"),
        ("인증 헤더 규약", "Authorization: Bearer <access_token>", "로그인 시 발급되는 JWT Access Token을 모든 보호된 엔드포인트 헤더에 주입"),
        ("증빙 파일 서빙 규약", "/uploads/{file_path}?ticket=<download_ticket>", "정적 파일 직접 노출 차단. 1회용 단일 소비(Single-use) 티켓으로만 파일 접근 가능"),
    ]),
    ("2. 공통 응답 포맷 (SrFormat)", [
        ("성공 응답 구조", "{\n  \"success\": true,\n  \"data\": <ResultData>,\n  \"message\": \"성공 메시지\"\n}", "모든 API 응답은 SrFormat 래퍼로 감싸져 반환되므로 FE에서는 res.data를 추출하여 사용"),
        ("에러 응답 구조 (4xx, 5xx)", "{\n  \"success\": false,\n  \"detail\": \"오류 원인 안내 메시지\",\n  \"error_code\": \"ERR_...\" (선택)\n}", "FastAPI 표준 HTTPException 및 유효성 검증 실패 시 detail 문자열 또는 상세 객체 전달"),
        ("동시성 충돌 (409 Conflict)", "HTTP 409 Conflict 반환", "승인 처리 중 학생이 삭제를 시도하거나, 이미 심사 완료된 증빙을 중복 처리 시 발생 (FE에서 즉각 갱신 안내 필요)"),
    ]),
    ("3. 핵심 3대 식별자 체계 (ID Conventions)", [
        ("로그인 계정 ID (user_id)", "users.id (VARCHAR/TEXT)", "로그인 시 사용하는 계정 아이디 (예: 'admin', 'teacher1', 'student1')"),
        ("사용자 고유 UUID (user_uuid)", "users.uuid (UUID/VARCHAR)", "시스템 내부 보안 식별용 UUID"),
        ("학적/교원 고유 PK (domain_id)", "students.student_id / teachers.teachers_id (INT)", "외래키 및 성적/상벌점 통계의 기준이 되는 정수형 Primary Key (예: studentId 101, 102...)"),
    ]),
    ("4. 주요 계정 안내 (초기 시드 데이터)", [
        ("최고 관리자 (Admin)", "ID: admin / PW: admin1234!", "전교생 현황, 학년도 활성화/복제, 33개 기준 배점 CRUD, 통계 내보내기, 감사로그, AI 감사"),
        ("교감 (Vice Principal)", "ID: vp / PW: password123!", "전교생 열람 및 직권 심사 권한, 전체 통계 열람"),
        ("1학년 부장교사 (Grade Head)", "ID: head1 / PW: password123!", "1학년 전 학급(1-1, 1-2...) 학생 목록, 심사, 상벌점 관리"),
        ("1-1반 담임교사 (Homeroom)", "ID: teacher1 / PW: password123!", "1학년 1반 담당 학생 전담 심사 및 상벌점 관리, 학급 AI 상담"),
        ("1-2반 담임교사 (Homeroom)", "ID: teacher2 / PW: password123!", "1학년 2반 담당 학생 전담 심사 및 상벌점 관리"),
        ("학생 계정 (Students)", "ID: student1 ~ student16 / PW: password123!", "student1~8 (1-1반), student9~16 (1-2반). 5대 역량 제출, 상벌점, AI 코칭"),
    ])
]

row_idx = 4
for sec_title, items in overview_sections:
    ws1.cell(row=row_idx, column=1, value=sec_title).font = Font(name="Pretendard", size=11, bold=True, color="0F172A")
    ws1.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=3)
    ws1.cell(row=row_idx, column=1).fill = PatternFill(start_color=C_SECTION_BG, end_color=C_SECTION_BG, fill_type="solid")
    row_idx += 1
    
    ws1.cell(row=row_idx, column=1, value="구분 / 항목").font = sub_header_font
    ws1.cell(row=row_idx, column=1).fill = sub_header_fill
    ws1.cell(row=row_idx, column=2, value="명세 및 형식").font = sub_header_font
    ws1.cell(row=row_idx, column=2).fill = sub_header_fill
    ws1.cell(row=row_idx, column=3, value="상세 설명 / FE 가이드").font = sub_header_font
    ws1.cell(row=row_idx, column=3).fill = sub_header_fill
    row_idx += 1
    
    s_r = row_idx
    for k, v, desc in items:
        c1 = ws1.cell(row=row_idx, column=1, value=k)
        c2 = ws1.cell(row=row_idx, column=2, value=v)
        c3 = ws1.cell(row=row_idx, column=3, value=desc)
        c1.font = data_font_bold
        c2.font = data_font_code
        c3.font = data_font
        c2.alignment = Alignment(wrap_text=True, vertical="center")
        c3.alignment = Alignment(wrap_text=True, vertical="center")
        row_idx += 1
    style_table(ws1, 1, 3, s_r, row_idx - 1)
    row_idx += 1

auto_width(ws1, {"A": 30, "B": 50, "C": 70})

# ==============================================================================
# SHEET 2: 02_전체_API_명세 (API_Directory)
# ==============================================================================
ws2 = wb.create_sheet(title="02_전체_API_명세")
ws2.views.sheetView[0].showGridLines = True

ws2.append(["백엔드 전체 API 명세 디렉터리 (Total 45 Core Endpoints)"])
ws2.append(["FastAPI 기반 REST API 엔드포인트, 인가 권한, 요청/응답 파라미터 및 프론트엔드 연동 주의사항"])
ws2.append([])

ws2["A1"].font = title_font
ws2["A2"].font = subtitle_font

api_headers = [
    "No", "도메인", "기능명", "Method", "Endpoint Path", "인가 권한 (RBAC)", 
    "요청 데이터 형식", "주요 Request 파라미터 / DTO", "주요 Response 데이터 필드", "프론트엔드 연동 및 예외처리 가이드"
]
ws2.append(api_headers)
for col_idx in range(1, len(api_headers) + 1):
    c = ws2.cell(row=4, column=col_idx)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

apis = [
    # 1. 인증/계정
    ("1", "인증", "사용자 로그인", "POST", "/api/auth/login", "누구나 (공개)", "JSON Body", 
     "{ id, password }", "{ access_token, token_type: 'bearer', user: { id, role, name, grade, class_no, number } }", 
     "로그인 성공 시 JWT 토큰 메모리(TokenManager) 및 세션에 저장. 5회 불일치 시 계정 잠금"),
    ("2", "인증", "사용자 로그아웃", "POST", "/api/auth/logout", "인증 사용자", "Empty / Token", 
     "헤더 Authorization: Bearer", "{ success: true, message: '로그아웃되었습니다.' }", 
     "서버 Redis 블랙리스트 등록 및 클라이언트 인메모리 토큰 파기. 로그인 화면으로 이동"),
    ("3", "인증", "내 정보 조회 (세션 확인)", "GET", "/api/auth/me", "인증 사용자", "Query / Header", 
     "헤더 Authorization: Bearer", "{ id, role, name, email, grade, class_no, number, studentId, teacherId }", 
     "새로고침 또는 앱 진입 시 세션 복원 및 사용자 권한/학급 정보 동기화"),
    ("4", "인증", "회원가입", "POST", "/api/auth/register", "누구나 (공개)", "JSON Body", 
     "{ id, password, name, role: 'student'|'teacher', email, otp_code, grade?, class_no?, number? }", 
     "{ user: { id, name, role } }", 
     "이메일 인증(OTP) 완료 후 가입 가능. 유효성 검증 실패 시 400 반환"),
    ("5", "인증", "이메일 OTP 인증번호 발송", "POST", "/api/auth/send-otp", "누구나 (공개)", "JSON Body", 
     "{ email, purpose: 'register'|'email_change' }", "{ success: true, message: '인증번호가 발송되었습니다.' }", 
     "부산SW마이스터고 이메일 도메인 검증 및 레이트 리밋(60초 재전송 대기 UI 필요)"),
    ("6", "인증", "이메일 OTP 인증번호 검증", "POST", "/api/auth/verify-otp", "누구나 (공개)", "JSON Body", 
     "{ email, otp_code, purpose }", "{ verified: true }", 
     "OTP 6자리 검증 성공 시 가입/변경 폼의 '인증완료' 상태 활성화"),
    ("7", "인증", "이메일 변경", "PATCH", "/api/auth/email", "인증 사용자", "JSON Body", 
     "{ new_email, otp_code, current_password }", "{ email: new_email }", 
     "현재 비밀번호 확인과 신규 이메일 OTP 검증을 동시에 수행하는 2단계 보안 엔드포인트"),

    # 2. 학생 도메인
    ("8", "학생", "5대 역량 인증 현황 조회", "GET", "/api/students/{student_id}/certification-status", "학생 본인 / 담당교사 / 관리자", "Path Param", 
     "student_id (INT 또는 'me')", "{ cert_status: '인증'|'미인증', total_score, areas: [{ area, score, maxScore, certified, items: [...] }] }", 
     "학생 대시보드 메인 렌더링용. 5대 영역 취득점수 및 게이지바 표시에 활용"),
    ("9", "학생", "증빙자료 신규 제출", "POST", "/api/submissions", "학생 본인", "multipart/form-data", 
     "year, area, itemId, activityDate, detail, description, link(선택), file(선택, 50MB 한도)", 
     "{ id, studentId, area, itemId, itemName, status: '제출완료', submittedAt }", 
     "50MB 클라이언트 사전 차단 필수. 링크 입력 시 http/https 프로토콜 검증"),
    ("10", "학생", "내 증빙자료 제출 내역 조회", "GET", "/api/submissions", "학생 본인", "Query Param", 
     "year?, area?, status?", "[{ id, year, area, itemId, itemName, status, score, teacherComment, submittedAt, reviewedAt, fileName, link }]", 
     "제출완료/인정완료/반려/재제출요청 상태별 탭 필터링 및 1회용 다운로드 티켓 연계"),
    ("11", "학생", "증빙자료 단건 상세 조회", "GET", "/api/submissions/{submission_id}", "학생 본인 / 심사교사 / 관리자", "Path Param", 
     "submission_id", "{ id, studentId, studentName, year, area, itemId, itemName, detail, description, status, score, maxScore, teacherComment, ... }", 
     "제출 서류 상세 열람 및 교사 심사 상세 화면의 단일 소스"),
    ("12", "학생", "증빙자료 보완 재제출", "PATCH", "/api/submissions/{submission_id}/resubmit", "학생 본인", "multipart/form-data", 
     "detail?, description?, link?, file?(50MB)", "{ id, status: '제출완료', resubmittedAt }", 
     "상태가 '반려' 또는 '재제출요청'인 건만 수정 가능. 인정완료 건은 403 차단"),
    ("13", "학생", "증빙자료 취소/삭제", "DELETE", "/api/submissions/{submission_id}", "학생 본인", "Path Param", 
     "submission_id", "{ success: true, message: '증빙자료가 삭제되었습니다.' }", 
     "심사 대기('제출완료') 상태만 삭제 가능. 교사가 이미 승인 중인 경우 409 동시성 에러 반환"),
    ("14", "학생", "증빙 파일 다운로드 티켓 발급", "POST", "/api/submissions/{submission_id}/download-ticket", "학생 본인 / 담당교사 / 관리자", "Path Param", 
     "submission_id", "{ ticket: '<UUID_TICKET>', download_url: '/uploads/...?ticket=...' }", 
     "IDOR 방어 1회용 티켓 발급(유효시간 3분). 발급받은 티켓 URL로 새 탭 또는 다운로드 유도"),
    ("15", "학생", "내 상벌점 이력 조회", "GET", "/api/points", "학생 본인", "Query Param", 
     "year?", "{ net_points, total_merit, total_demerit, items: [{ point_id, year, type: '상점'|'벌점', points, reason, assigned_date }] }", 
     "벌점은 음수 부호로 처리되어 net_points 계산. 상벌점 내역 테이블 렌더링"),

    # 3. 교사 도메인
    ("16", "교사", "교사 대시보드 통계 조회", "GET", "/api/teacher/dashboard", "교사 / 관리자", "None", 
     "헤더 Authorization: Bearer", "{ scope: { role, grade, classNo }, pending_count, total_students, certified_students, recent_submissions: [...] }", 
     "담임/부장/교감 교무분장 직책에 따라 통계 스코프가 자동 분기되어 반환됨"),
    ("17", "교사", "담당 학생 목록 및 현황 검색", "GET", "/api/teacher/students", "교사 / 관리자", "Query Param", 
     "grade?, class_no?, search?, cert_status?", "[{ student_id, user_id, name, grade, class_no, number, total_score, cert_status, pending_count }]", 
     "담임은 자기 반만 조회, 학년부장은 해당 학년, 교감/관리자는 전교생 조회 허용"),
    ("18", "교사", "특정 학생 상세 및 제출 목록", "GET", "/api/teacher/students/{student_id}", "담당 교사 / 관리자", "Path Param", 
     "student_id", "{ student: {...}, certification_status: {...}, submissions: [...], points: [...] }", 
     "타 반 학생 접근 시 403 Forbidden 차단(교무분장 인가 매트릭스 엄격 검증)"),
    ("19", "교사", "증빙자료 심사 승인", "PATCH", "/api/submissions/{submission_id}/approve", "담당 교사 / 관리자", "JSON Body", 
     "{ score: INT/FLOAT, teacherComment?: STRING }", "{ id, status: '인정완료', score, reviewerId, studentTotals: {...} }", 
     "승인 성공 시 studentTotals가 함께 반환되므로 FE에서 학생 총점 및 게이지 즉시 갱신"),
    ("20", "교사", "증빙자료 심사 반려/재제출요청", "PATCH", "/api/submissions/{submission_id}/reject", "담당 교사 / 관리자", "JSON Body", 
     "{ reject_status: '반려'|'재제출요청', teacherComment: STRING(필수) }", "{ id, status: '반려'|'재제출요청', teacherComment }", 
     "반려 사유(teacherComment) 필수 입력. 사유 누락 시 422/400 유효성 에러"),
    ("21", "교사", "기승인 건 인정 점수 사후 수정", "PATCH", "/api/submissions/{submission_id}/score", "담당 교사 / 관리자", "JSON Body", 
     "{ score: INT/FLOAT, teacherComment: STRING(수정사유) }", "{ id, status: '인정완료', score, scoreHistory: [...] }", 
     "이미 인정완료된 건의 점수를 재조정할 때 호출. 항목 최대 배점(maxScore) 초과 불가"),
    ("22", "교사", "상벌점 신규 부여", "POST", "/api/points", "교사 / 관리자", "JSON Body", 
     "{ student_id: INT, year: INT, type: '상점'|'벌점', points: INT(벌점시 음수), reason: STRING, assigned_date: 'YYYY-MM-DD' }", 
     "{ point_id, student_id, net_points }", 
     "벌점 부여 시 반드시 음수 부호(-3) 보존 전송. student_id는 문자열이 아닌 정수 PK 매핑 필수"),
    ("23", "교사", "상벌점 내역 수정", "PATCH", "/api/points/{point_id}", "부여 교사 / 관리자", "JSON Body", 
     "{ points?: INT, reason?: STRING, assigned_date?: STRING }", "{ point_id, updated_fields }", 
     "타 교사가 부여한 건은 담임/부장/관리자 외 수정 불가"),
    ("24", "교사", "상벌점 내역 삭제 (소프트 삭제)", "DELETE", "/api/points/{point_id}", "부여 교사 / 관리자", "Path Param", 
     "point_id", "{ success: true, message: '상벌점이 삭제되었습니다.' }", 
     "데이터베이스에서 물리 삭제되지 않고 deleted_at 처리되어 감사 이력 보존"),

    # 4. 관리자 도메인
    ("25", "관리자", "전교생 인증 현황 KPI 대시보드", "GET", "/api/admin/overview", "관리자 전용", "Query Param", 
     "year?", "{ kpis: { total_students, certified_count, cert_rate, pending_submissions }, area_rates: [...], grade_stats: [...] }", 
     "전교생 종합 지표 시각화 및 통계 차트 데이터 소스"),
    ("26", "관리자", "전체 학년도 목록 조회", "GET", "/api/years", "누구나 (인증 권장)", "None", 
     "None", "[{ year: 2026, is_active: true, created_at, criteria_count }]", 
     "시스템 전체에서 사용 중인 기준 학년도 목록 반환. 활성 학년도(is_active: true) 우선 선택"),
    ("27", "관리자", "신규 학년도 등록", "POST", "/api/years", "관리자 전용", "JSON Body", 
     "{ year: INT }", "{ year, is_active: false, message: '등록 완료' }", 
     "신규 학년도 등록 시 기본적으로 비활성 상태로 생성"),
    ("28", "관리자", "학년도 활성화/비활성화 토글", "PATCH", "/api/years/{year}", "관리자 전용", "JSON Body", 
     "{ is_active: BOOLEAN }", "{ year, is_active }", 
     "활성화 변경 시 즉시 전교생 증빙 제출 및 심사의 기준 학년도로 적용됨"),
    ("29", "관리자", "이전 학년도 기준 복제", "POST", "/api/years/{target_year}/copy-from/{source_year}", "관리자 전용", "Path Param", 
     "target_year, source_year", "{ target_year, copied_areas_count, copied_items_count: 33 }", 
     "새 학년도 시작 시 이전 학년도의 5대 영역 및 33개 세부항목 배점 설정을 통째로 복제"),
    ("30", "관리자", "평가 기준 및 배점 조회", "GET", "/api/admin/criteria", "관리자 전용", "Query Param", 
     "year?", "[{ grade: 1, areas: [{ area_id, area, max_score, items: [{ item_id, name, max_score, scoring_type, ... }] }] }]", 
     "학년도별 1, 2, 3학년 5대 영역 및 33개 세부항목 전체 트리 반환"),
    ("31", "관리자", "평가 영역 배점 수정", "PATCH", "/api/admin/criteria/areas/{area_id}", "관리자 전용", "JSON Body", 
     "{ max_score: INT }", "{ area_id, max_score }", 
     "영역별 상한 배점 수정(예: 전공역량 25점 -> 30점). 실패 시 400 반환"),
    ("32", "관리자", "세부 평가 항목 신규 추가", "POST", "/api/admin/criteria/areas/{area_id}/items", "관리자 전용", "JSON Body", 
     "{ name, max_score, scoring_type: '고정'|'누적'|'최고', requires_evidence: BOOL }", 
     "{ item_id, area_id, name, max_score }", 
     "영역 내에 새로운 평가 항목을 생성할 때 사용"),
    ("33", "관리자", "세부 평가 항목 수정", "PATCH", "/api/admin/criteria/items/{item_id}", "관리자 전용", "JSON Body", 
     "{ name?, max_score?, scoring_type?, requires_evidence?, active?: BOOL }", 
     "{ item_id, updated_fields }", 
     "배점 및 활성화 여부(active: false 시 신규 제출 목록에서 숨김) 수정"),
    ("34", "관리자", "세부 평가 항목 삭제", "DELETE", "/api/admin/criteria/items/{item_id}", "관리자 전용", "Path Param", 
     "item_id", "{ success: true, message: '항목이 삭제되었습니다.' }", 
     "소프트 삭제(is_active: false) 처리되어 기존 학생들의 기취득 이력은 안전하게 보존"),
    ("35", "관리자", "학생 명단 엑셀/CSV 일괄 등록", "POST", "/api/admin/students/batch", "관리자 전용", "multipart/form-data", 
     "file: EXCEL/CSV", "{ total_count, success_count, failed_count, errors: [{ row, reason }] }", 
     "신학기 전교생 일괄 등록. 학번, 이름, 이메일, 초기 비밀번호 자동 생성"),
    ("36", "관리자", "NEIS 표준 양식 학생 명단 일괄 등록", "POST", "/api/admin/students/batch/neis", "관리자 전용", "multipart/form-data", 
     "file: EXCEL/CSV (NEIS 규격)", "{ imported_count, class_mappings }", 
     "나이스(NEIS) 학적 변동 양식과 100% 호환되는 자동 파싱 파이프라인"),
    ("37", "관리자", "교원 명단 엑셀/CSV 일괄 등록", "POST", "/api/admin/teachers/batch", "관리자 전용", "multipart/form-data", 
     "file: EXCEL/CSV", "{ total_teachers, assigned_classes }", 
     "교사 계정 생성 및 담임/부장/교감 교무분장 일괄 배정"),
    ("38", "관리자", "교내 통계 내보내기", "GET", "/api/admin/export/stats", "관리자 전용", "Query Param", 
     "format: 'csv'|'xlsx'|'pdf', year?", "바이너리 파일 스트림 (FileResponse)", 
     "응답 Content-Type에 맞춰 브라우저에서 자동 다운로드 저장되도록 blob 처리"),
    ("39", "관리자", "시스템 종합 감사 로그 조회", "GET", "/api/admin/audit-logs", "관리자 전용", "Query Param", 
     "page?, limit?, action?, user_id?, start_date?, end_date?", 
     "{ total, logs: [{ log_id, user_id, action, ip_address, created_at, details }] }", 
     "심사 승인/반려, 점수 수정, 상벌점 변동, 기준 변경 등 시스템 내 모든 행위 감사 추적"),
    ("40", "관리자", "데이터 수동 재계산 배치", "POST", "/api/admin/recalculate", "관리자 전용", "JSON Body", 
     "{ year?: INT }", "{ processed_students, recalculated_scores, duration_ms }", 
     "평가 기준이나 배점이 변경되었을 때 기존 학생들의 취득 점수 및 인증 판정 일괄 재계산"),
    ("41", "관리자", "고아 파일 정리 가비지 컬렉터(GC)", "POST", "/api/admin/storage/gc", "관리자 전용", "None", 
     "None", "{ scanned_files, orphan_deleted, freed_bytes }", 
     "제출 취소되거나 삭제되어 DB 연결이 끊긴 업로드 미디어 파일 물리적 영구 삭제"),

    # 5. AI 서비스
    ("42", "AI", "증빙자료 자동 분류 및 항목 추천", "POST", "/api/ai/classify-submission", "학생 / 교사 / 관리자", "JSON Body", 
     "{ detail: STRING, description: STRING }", "{ area: '전공역량', itemName: '알고리즘 경진대회', confidence: 0.94, reasoning: STRING }", 
     "학생이 증빙 활동내용 작성 시 디바운스(600ms)로 호출하여 영역/항목 자동 선택 추천"),
    ("43", "AI", "학생 맞춤형 역량 달성 도우미 챗봇", "POST", "/api/ai/student/chat", "학생 전용", "JSON Body", 
     "{ message: STRING, history?: [{ role, content }] }", "{ reply: STRING, recommended_actions: [...] }", 
     "학생의 현재 5대 역량 취득 현황과 부족한 영역을 분석하여 맞춤형 활동 및 자격증 추천"),
    ("44", "AI", "교사용 심사 검토의견 초안 생성", "POST", "/api/ai/teacher/review-draft", "교사 / 관리자", "JSON Body", 
     "{ submission_id: INT }", "{ draft_comment: STRING, suggested_score: INT, confidence: FLOAT, risk_flags: [...] }", 
     "교사 심사 화면에서 'AI 초안 생성' 버튼 클릭 시 3초 내로 모범 검토의견 및 권장 점수 제안"),
    ("45", "AI", "부정 의심 제출 탐지 및 이상치 감사", "POST", "/api/ai/admin/audit-anomalies", "관리자 전용", "JSON Body", 
     "{ year?: INT }", "{ anomaly_count, high_risk_submissions: [{ submission_id, student_name, reason, risk_score }] }", 
     "동일 증빙 다중 제출, 비정상적 단기간 점수 폭증, 위조 의심 링크/파일 탐지 및 감사 리포트")
]

row_start = 5
for item in apis:
    r = ws2.max_row + 1
    ws2.append(list(item))
    
    # Method cell highlight
    m_cell = ws2.cell(row=r, column=4)
    m_cell.font = Font(name="Consolas", size=9.5, bold=True)
    m_cell.alignment = Alignment(horizontal="center", vertical="center")
    m_val = str(item[3]).upper()
    if m_val == "GET":
        m_cell.fill = PatternFill(start_color=C_BLUE_BG, end_color=C_BLUE_BG, fill_type="solid")
        m_cell.font = Font(name="Consolas", size=9.5, bold=True, color=C_BLUE_TXT)
    elif m_val == "POST":
        m_cell.fill = PatternFill(start_color=C_GREEN_BG, end_color=C_GREEN_BG, fill_type="solid")
        m_cell.font = Font(name="Consolas", size=9.5, bold=True, color=C_GREEN_TXT)
    elif m_val in ("PATCH", "PUT"):
        m_cell.fill = PatternFill(start_color=C_AMBER_BG, end_color=C_AMBER_BG, fill_type="solid")
        m_cell.font = Font(name="Consolas", size=9.5, bold=True, color=C_AMBER_TXT)
    elif m_val == "DELETE":
        m_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        m_cell.font = Font(name="Consolas", size=9.5, bold=True, color=C_ACCENT_RED)
        
    ws2.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws2.cell(row=r, column=2).alignment = Alignment(horizontal="center", vertical="center")
    ws2.cell(row=r, column=5).font = data_font_code
    ws2.cell(row=r, column=6).alignment = Alignment(horizontal="center", vertical="center")
    ws2.cell(row=r, column=7).alignment = Alignment(horizontal="center", vertical="center")
    ws2.cell(row=r, column=8).font = data_font_code
    ws2.cell(row=r, column=8).alignment = Alignment(wrap_text=True, vertical="center")
    ws2.cell(row=r, column=9).font = data_font_code
    ws2.cell(row=r, column=9).alignment = Alignment(wrap_text=True, vertical="center")
    ws2.cell(row=r, column=10).font = data_font
    ws2.cell(row=r, column=10).alignment = Alignment(wrap_text=True, vertical="center")

style_table(ws2, 1, 10, row_start, ws2.max_row)
auto_width(ws2, {"A": 6, "B": 10, "C": 24, "D": 10, "E": 38, "F": 20, "G": 18, "H": 35, "I": 35, "J": 45})

# ==============================================================================
# SHEET 3: 03_FE_화면별_연동가이드 (FE_Checklist)
# ==============================================================================
ws3 = wb.create_sheet(title="03_FE_화면별_연동가이드")
ws3.views.sheetView[0].showGridLines = True

ws3.append(["프론트엔드(FE) 주요 화면별 연동 체크리스트 & UI 예외처리 가이드"])
ws3.append(["각 화면에서 필수 호출하는 API, 필수 상태(로딩/에러/성공), UX 인터랙션 처리 지침"])
ws3.append([])

ws3["A1"].font = title_font
ws3["A2"].font = subtitle_font

fe_headers = ["화면 ID", "역할", "화면명 / URL 경로", "연동 핵심 API", "필수 입력 및 유효성 검증", "권장 UX 인터랙션 & 상태 피드백", "백엔드 예외 및 에러 처리 지침"]
ws3.append(fe_headers)
for col_idx in range(1, len(fe_headers) + 1):
    c = ws3.cell(row=4, column=col_idx)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

fe_screens = [
    ("SCR-01", "공통", "로그인 화면\n(/login)", "POST /api/auth/login", 
     "ID/PW 필수 입력, 공백 제거", "로그인 진행 중 '인증 중...' 버튼 비활성화, 빠른 테스트용 역할별 스위처 제공", 
     "401 시 '아이디 또는 비밀번호가 일치하지 않습니다' 노출. 5회 연속 실패 시 계정 잠금 안내"),
    ("SCR-02", "공통", "회원가입 화면\n(/signup)", "POST /api/auth/send-otp\nPOST /api/auth/verify-otp\nPOST /api/auth/register", 
     "학교 이메일(@bssm.hs.kr), 비밀번호 8자 이상(영문+숫자+특수문자), OTP 6자리", 
     "OTP 전송 후 60초 타이머 동작, 이메일 인증 완료 전에는 '가입하기' 버튼 비활성화", 
     "이메일 중복 시 400 반환, 인증번호 불일치 시 인라인 에러 빨간색 표기"),
    ("SCR-03", "학생", "학생 대시보드\n(/student/dashboard)", "GET /api/auth/me\nGET /api/students/{id}/certification-status\nGET /api/years", 
     "로그인 세션 확인", "5대 역량 취득 현황 게이지 바(취득점수/상한점수), 인증 판정 배지('인증완료' vs '미인증'), 부족 역량 강조", 
     "활성 학년도 없을 시 관리자 문의 배너 표출, API 지연 시 스켈레톤 로더 적용"),
    ("SCR-04", "학생", "증빙자료 제출\n(/student/submit)", "GET /api/admin/criteria\nPOST /api/ai/classify-submission\nPOST /api/submissions", 
     "활동일자, 영역/세부항목 선택, 활동상세, 설명, 파일 또는 링크 중 1개 필수", 
     "활동내용 입력 시 AI 자동분류 추천 배지 클릭 시 항목 자동 세팅. 파일 첨부 시 50MB 용량 클라이언트 사전 체크", 
     "50MB 초과 시 파일 즉시 거부 및 경고, 전송 중 프로그레스 스피너 및 중복 클릭 차단"),
    ("SCR-05", "학생", "제출 내역 및 보완\n(/student/submissions)", "GET /api/submissions\nPATCH /api/submissions/{id}/resubmit\nDELETE /api/submissions/{id}\nPOST /api/submissions/{id}/download-ticket", 
     "재제출 시 보완 설명 및 수정 파일/링크", "상태별(전체/제출완료/인정완료/반려) 탭 필터링, 증빙 파일 클릭 시 1회용 다운로드 티켓 자동 발급 후 새 탭 열기", 
     "교사가 심사 중 삭제 시도 시 409 Conflict 감지 -> '이미 심사 중인 건입니다' 토스트 띄우고 목록 자동 새로고침"),
    ("SCR-06", "학생", "상벌점 내역\n(/student/points)", "GET /api/points", 
     "연도별 조회", "총 상점, 총 벌점, 최종 상벌점(Net Points) 요약 카드, 상점(초록) 및 벌점(빨강) 시각적 대비", 
     "벌점이 0점 미만일 때 최종 합산 점수 클램핑(음수 처리) 규약 정상 표시"),
    ("SCR-07", "학생", "AI 역량상담 챗봇\n(/student/chat)", "POST /api/ai/student/chat", 
     "상담 메시지 (최대 500자)", "말풍선 대화 UI, 추천 활동 바로가기 버튼 칩, 전송 중 타이핑 애니메이션", 
     "AI API 일시 지연 또는 429 할당량 초과 시 '내장 룰 기반 추천 가이드' 자동 폴백 메시지 표출"),
    ("SCR-08", "교사", "교사 대시보드\n(/teacher/dashboard)", "GET /api/teacher/dashboard", 
     "교사 권한 확인", "담임/부장/교감 권한 배지 표기, '심사 대기 건수' 알림 배지, 최근 제출된 증빙 카드 리스트", 
     "담당 학급이 미배정된 교사의 경우 '담당 학급 배정 대기 중' 안내 및 읽기 전용 모드 전환"),
    ("SCR-09", "교사", "담당 학생 관리\n(/teacher/students)", "GET /api/teacher/students", 
     "학생명 검색어, 학년/반 필터", "학생 테이블(학번, 이름, 취득점수, 인증여부, 미심사 건수), 행 클릭 시 학생 상세 페이지 이동", 
     "교무분장 인가 매트릭스에 따라 타 반 학생 검색 시 403 차단 및 비노출"),
    ("SCR-10", "교사", "증빙 심사 및 점수수정\n(/teacher/review/[id])", "GET /api/submissions/{id}\nPOST /api/ai/teacher/review-draft\nPATCH /api/submissions/{id}/approve\nPATCH /api/submissions/{id}/reject\nPATCH /api/submissions/{id}/score", 
     "승인 시 점수(최대배점 이하), 반려 시 반려사유(필수), 점수수정 시 수정사유(필수)", 
     "'AI 심사 초안 생성' 원클릭으로 검토의견 자동 입력. 승인 완료 건은 '사후 점수 수정' 폼으로 자동 전환", 
     "심사 완료 즉시 studentTotals 응답을 반영하여 화면 새로고침 없이 최신 점수 및 상태 렌더링"),
    ("SCR-11", "교사", "상벌점 관리\n(/teacher/points)", "POST /api/points\nPATCH /api/points/{id}\nDELETE /api/points/{id}", 
     "대상 학생(학번/이름), 구분(상점/벌점), 점수(1~5), 사유, 일자", 
     "벌점 선택 시 자동으로 음수 부호(-) 표시, 등록 즉시 목록 반영, 부여 교사 본인만 수정/삭제 가능", 
     "학생 ID 전달 시 문자열(예: 'student1')이 아닌 고유 정수 PK(studentId) 매핑 전송"),
    ("SCR-12", "관리자", "전교생 현황 개요\n(/admin/overview)", "GET /api/admin/overview\nGET /api/admin/export/stats", 
     "학년도 선택", "전교생 인증률 KPI, 5대 영역별 평균 달성률 막대 차트, 엑셀/CSV/PDF 다운로드 버튼", 
     "PDF/엑셀 다운로드 클릭 시 브라우저 블롭(Blob) 다운로드 스트림 처리"),
    ("SCR-13", "관리자", "기준 및 배점 관리\n(/admin/criteria)", "GET /api/admin/criteria\nPATCH /api/admin/criteria/areas/{id}\nPATCH /api/admin/criteria/items/{id}\nPOST /api/admin/criteria/areas/{id}/items\nDELETE /api/admin/criteria/items/{id}", 
     "영역 배점, 항목명, 항목 배점(1~50), 산정방식, 증빙필수여부", 
     "학년도/학년별 탭 선택, 인라인 배점 편집(수정 즉시 저장), 세부 항목 추가/삭제 모달", 
     "배점 수정 실패 시 빈 catch로 삼키지 않고 상단 에러 배너 노출 및 이전 값으로 롤백"),
    ("SCR-14", "관리자", "학년도 및 기준 복제\n(/admin/years)", "GET /api/years\nPOST /api/years\nPATCH /api/years/{year}\nPOST /api/years/{target}/copy-from/{source}", 
     "신규 연도(4자리 숫자), 복제 대상 연도", "현재 활성 학년도 스위치 토글, '이전 학년도 기준 복제' 확인 다이얼로그", 
     "중복 학년도 생성 시 400 에러 처리, 비활성화 토글 시 즉시 상태 반영"),
    ("SCR-15", "관리자", "명단 일괄 등록 & NEIS\n(/admin/batch)", "POST /api/admin/students/batch\nPOST /api/admin/students/batch/neis\nPOST /api/admin/teachers/batch", 
     "엑셀(.xlsx) 또는 CSV 파일 업로드", "샘플 서식 다운로드 링크 제공, 업로드 후 성공 행수 / 실패 행별 오류 사유 리포트 모달 표시", 
     "엑셀 내 필수 컬럼(학번, 이름) 누락 시 해당 행 번호와 실패 원인을 테이블로 안내"),
    ("SCR-16", "관리자", "감사로그 & AI 이상탐지\n(/admin/audit)", "GET /api/admin/audit-logs\nPOST /api/ai/admin/audit-anomalies\nPOST /api/admin/recalculate", 
     "필터(행위자, 액션, 기간)", "시간 역순 감사 로그 타임라인, '부정 의심 제출 AI 탐지' 카드(위험 점수 80점 이상 하이라이트)", 
     "데이터 정합성 이상 발생 시 '전체 데이터 재계산(Recalculate)' 버튼 제공")
]

row_start = 5
for item in fe_screens:
    r = ws3.max_row + 1
    ws3.append(list(item))
    ws3.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws3.cell(row=r, column=1).font = data_font_bold
    ws3.cell(row=r, column=2).alignment = Alignment(horizontal="center", vertical="center")
    ws3.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="center")
    ws3.cell(row=r, column=3).font = data_font_bold
    ws3.cell(row=r, column=4).font = data_font_code
    ws3.cell(row=r, column=4).alignment = Alignment(wrap_text=True, vertical="center")
    ws3.cell(row=r, column=5).alignment = Alignment(wrap_text=True, vertical="center")
    ws3.cell(row=r, column=6).alignment = Alignment(wrap_text=True, vertical="center")
    ws3.cell(row=r, column=7).alignment = Alignment(wrap_text=True, vertical="center")

style_table(ws3, 1, 7, row_start, ws3.max_row)
auto_width(ws3, {"A": 12, "B": 10, "C": 24, "D": 38, "E": 32, "F": 42, "G": 42})

# ==============================================================================
# SHEET 4: 04_교무분장_권한매트릭스 (Role_Matrix)
# ==============================================================================
ws4 = wb.create_sheet(title="04_교무분장_권한매트릭스")
ws4.views.sheetView[0].showGridLines = True

ws4.append(["교무분장 및 역할별 차등 인가 매트릭스 (Role-Based Access Control)"])
ws4.append(["학교 현장의 8대 직책/역할별 기능 접근 권한 및 심사/상벌점 부여 가능 범위"])
ws4.append([])

ws4["A1"].font = title_font
ws4["A2"].font = subtitle_font

matrix_headers = [
    "구분", "기능 / 업무 행위", "학생 본인", "정담임 교사", "부/임시담임", "학년부장", 
    "비담임 교과교사", "1순위 전담평가자", "교감 / 총괄부서장", "최고 관리자 (Admin)"
]
ws4.append(matrix_headers)
for col_idx in range(1, len(matrix_headers) + 1):
    c = ws4.cell(row=4, column=col_idx)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

matrix_rows = [
    ("인증현황", "본인 5대 역량 인증현황 조회", "O (본인)", "O (담당학급)", "O (담당학급)", "O (해당학년)", "O (전교생 열람)", "O (전교생 열람)", "O (전교생)", "O (전교생)"),
    ("학생조회", "학생 목록 검색 및 상세 프로필", "X", "O (담당학급)", "O (담당학급)", "O (해당학년)", "O (검색 허용)", "O (검색 허용)", "O (전교생)", "O (전교생)"),
    ("증빙제출", "증빙자료 신규 제출 (파일/링크)", "O (본인)", "X", "X", "X", "X", "X", "X", "X"),
    ("증빙수정", "반려 건 보완 재제출", "O (본인)", "X", "X", "X", "X", "X", "X", "X"),
    ("증빙삭제", "미심사 제출건 취소/삭제", "O (심사전)", "X", "X", "X", "X", "X", "X", "O (직권삭제)"),
    ("파일열람", "증빙 파일 다운로드 티켓 발급", "O (본인)", "O (담당학급)", "O (담당학급)", "O (해당학년)", "O (담당영역)", "O (전담영역)", "O (전교생)", "O (전교생)"),
    ("심사승인", "증빙자료 심사 승인 및 배점 부여", "X", "O (담당학급)", "O (담당학급)", "O (해당학년)", "△ (전담영역만)", "O (해당항목 전담)", "O (전교생 직권)", "O (전교생 직권)"),
    ("심사반려", "증빙자료 심사 반려 및 재제출요청", "X", "O (담당학급)", "O (담당학급)", "O (해당학년)", "△ (전담영역만)", "O (해당항목 전담)", "O (전교생 직권)", "O (전교생 직권)"),
    ("점수수정", "기승인 건 인정 점수 사후 수정", "X", "O (사유필수)", "O (사유필수)", "O (사유필수)", "X", "O (전담항목)", "O (직권수정)", "O (직권수정)"),
    ("상벌점부여", "상벌점 신규 등록 (벌점 음수)", "X", "O (담당학급)", "O (담당학급)", "O (해당학년)", "O (수업참여학생)", "O (관련학생)", "O (전교생)", "O (전교생)"),
    ("상벌점수정", "기부여 상벌점 수정 및 삭제", "X", "O (본인부여건)", "O (본인부여건)", "O (해당학년)", "O (본인부여건)", "O (본인부여건)", "O (전교생)", "O (전교생)"),
    ("대시보드", "교사 대시보드 스코프", "X", "담당학급 중심", "담당학급 중심", "해당학년 중심", "전체 요약", "전담영역 대기건", "전교생 총괄", "전교생 총괄"),
    ("기준관리", "학년도 생성 / 활성화 토글", "X", "X", "X", "X", "X", "X", "X", "O (관리자전용)"),
    ("배점설정", "33개 세부항목 배점 CRUD", "X", "X", "X", "X", "X", "X", "X", "O (관리자전용)"),
    ("일괄등록", "학생/교원 NEIS 엑셀 일괄 등록", "X", "X", "X", "X", "X", "X", "X", "O (관리자전용)"),
    ("통계내보내기", "전교생 통계 CSV/XLSX/PDF 출력", "X", "△ (학급명단)", "△ (학급명단)", "△ (학년명단)", "X", "X", "O (전교생)", "O (전교생)"),
    ("감사로그", "전체 시스템 감사 로그 조회", "X", "X", "X", "X", "X", "X", "X", "O (관리자전용)")
]

row_start = 5
for row in matrix_rows:
    r = ws4.max_row + 1
    ws4.append(list(row))
    ws4.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws4.cell(row=r, column=1).font = data_font_bold
    ws4.cell(row=r, column=2).alignment = Alignment(vertical="center")
    for c_idx in range(3, 11):
        cell = ws4.cell(row=r, column=c_idx)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        val = str(cell.value)
        if val.startswith("O"):
            cell.font = Font(name="Pretendard", size=9.5, bold=True, color=C_GREEN_TXT)
        elif val.startswith("X"):
            cell.font = Font(name="Pretendard", size=9.5, color="94A3B8")
        elif val.startswith("△"):
            cell.font = Font(name="Pretendard", size=9.5, color=C_AMBER_TXT)

style_table(ws4, 1, 10, row_start, ws4.max_row)
auto_width(ws4, {"A": 14, "B": 28, "C": 14, "D": 16, "E": 16, "F": 16, "G": 18, "H": 18, "I": 18, "J": 18})

# ==============================================================================
# SHEET 5: 05_5대역량_33개평가기준 (Criteria_33)
# ==============================================================================
ws5 = wb.create_sheet(title="05_5대역량_33개평가기준")
ws5.views.sheetView[0].showGridLines = True

ws5.append(["마이스터 역량인증제 5대 영역 및 33개 세부 평가 기준표 (2026학년도)"])
ws5.append(["영역별 상한 배점, 세부 항목명, 기본 배점, 점수 산정 방식 및 증빙자료 필수 여부"])
ws5.append([])

ws5["A1"].font = title_font
ws5["A2"].font = subtitle_font

crit_headers = ["영역 번호", "역량 영역명", "영역 상한배점", "항목 번호", "세부 평가 항목명", "기본 배점", "산정 방식 (Scoring Type)", "증빙 필수", "상세 설명 및 인정 기준"]
ws5.append(crit_headers)
for col_idx in range(1, len(crit_headers) + 1):
    c = ws5.cell(row=4, column=col_idx)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

criteria_33 = [
    # 1. 전공역량 (상한 25점)
    ("01", "전공역량", 25, "item-101", "전공 교과 성적 우수", 10, "누적합산형", "선택 (NEIS연동)", "전공 교과 성취도 A 또는 학기말 성적 우수자"),
    ("01", "전공역량", 25, "item-102", "전공 프로젝트 포트폴리오", 10, "누적합산형", "필수 (파일/깃허브)", "깃허브 레포지토리 링크 또는 소프트웨어 산출물 PDF"),
    ("01", "전공역량", 25, "item-103", "교내 전공 경진대회 수상", 8, "최고점수반영형", "필수 (상장 사본)", "해커톤, 알고리즘 대회, 임베디드 경진대회 입상"),
    ("01", "전공역량", 25, "item-104", "대외 전공 경진대회 입상", 10, "최고점수반영형", "필수 (상장 사본)", "전국 기능경기대회, 정부/공공기관 주최 IT 공모전 수상"),
    ("01", "전공역량", 25, "item-105", "전공 동아리 우수 활동", 5, "고정점수형", "필수 (활동일지)", "전공심화동아리 80% 이상 출석 및 결과물 발표"),
    ("01", "전공역량", 25, "item-106", "산학협력 프로젝트 참여", 8, "누적합산형", "필수 (수행확인서)", "기업 연계 산학맞춤 멘토링 및 실무 프로젝트 완수"),
    ("01", "전공역량", 25, "item-107", "오픈소스 기여 (Pull Request)", 5, "누적합산형", "필수 (PR 링크)", "공개 오픈소스 저장소 PR 머지 또는 기술 아티클 기고"),

    # 2. 외국어역량 (상한 20점)
    ("02", "외국어역량", 20, "item-201", "TOEIC 성적 우수 (700점 이상)", 15, "최고점수반영형", "필수 (성적표 사본)", "공인 토익 700점 이상(15점), 600점 이상(10점)"),
    ("02", "외국어역량", 20, "item-202", "TOEIC Bridge 우수 (140점 이상)", 10, "최고점수반영형", "필수 (성적표 사본)", "토익 브릿지 140점 이상 달성 시 취득"),
    ("02", "외국어역량", 20, "item-203", "교내 영어 말하기/어휘 대회", 8, "최고점수반영형", "필수 (상장 사본)", "교내 실무 영어 프레젠테이션 및 보카 경시대회 입상"),
    ("02", "외국어역량", 20, "item-204", "방과후 실무영어 성실 이수", 5, "고정점수형", "필수 (이수증)", "영어 방과후 학교 프로그램 80% 이상 출석 및 수료"),
    ("02", "외국어역량", 20, "item-205", "원어 기술 문서 번역/작성", 5, "누적합산형", "필수 (번역문/원본)", "영문 오픈소스 공식 문서 번역 및 기술 블로그 포스팅"),
    ("02", "외국어역량", 20, "item-206", "제2외국어 기초 회화 이수", 5, "고정점수형", "선택 (확인서)", "일본어/중국어 등 기초 외국어 과정 수료"),

    # 3. 자격취득 (상한 20점)
    ("03", "자격취득", 20, "item-301", "정보처리기능사 취득", 10, "고정점수형", "필수 (자격증 사본)", "한국산업인력공단 국가기술자격증 정보처리기능사"),
    ("03", "자격취득", 20, "item-302", "정보보안기능사 / 네트워크관리사", 10, "고정점수형", "필수 (자격증 사본)", "네트워크관리사 2급 또는 정보보안 분야 공인 자격증"),
    ("03", "자격취득", 20, "item-303", "SQLD / 리눅스마스터", 8, "고정점수형", "필수 (자격증 사본)", "SQL 개발자(SQLD) 또는 리눅스마스터 2급 이상"),
    ("03", "자격취득", 20, "item-304", "AWS / 클라우드 국제공인자격", 10, "최고점수반영형", "필수 (디지털 배지/자격증)", "AWS Certified Cloud Practitioner, SAA, GCP, Azure 등"),
    ("03", "자격취득", 20, "item-305", "컴퓨터활용능력 / 사무자격", 5, "최고점수반영형", "필수 (자격증 사본)", "컴퓨터활용능력 1급/2급, MOS Master 등"),
    ("03", "자격취득", 20, "item-306", "기타 공인 기술 자격증", 5, "누적합산형", "필수 (자격증 사본)", "교과협의회에서 인정한 전공 유관 IT 공인 자격증"),

    # 4. 인성/봉사 (상한 20점)
    ("04", "인성/봉사", 20, "item-401", "교내외 봉사활동 시간 (40시간 이상)", 10, "누적합산형", "필수 (VMS/1365)", "1365 또는 VMS 공인 봉사활동 시간 누적 증빙"),
    ("04", "인성/봉사", 20, "item-402", "학급 임원 및 학생회 간부 활동", 8, "고정점수형", "선택 (생활기록부)", "반장, 부반장, 학생회 부서장 활동 성실 수행"),
    ("04", "인성/봉사", 20, "item-403", "교내 선행 / 효행 / 모범 표창", 8, "누적합산형", "필수 (표창장)", "교내외 모범 학생 표창 및 선행상 수상"),
    ("04", "인성/봉사", 20, "item-404", "학교폭력 예방 및 또래상담 활동", 5, "고정점수형", "필수 (활동확인서)", "또래상담부 수료 및 캠페인 참여"),
    ("04", "인성/봉사", 20, "item-405", "상벌점 순상점 우수 (상점-벌점 >= 10)", 5, "고정점수형", "자동연동 (시스템)", "시스템 내 누적 순상점이 10점 이상인 경우 자동 가산"),
    ("04", "인성/봉사", 20, "item-406", "환경 정화 및 멘토링 봉사", 5, "누적합산형", "필수 (활동보고서)", "후배 전공 지도 멘토링 또는 지역사회 환경정화"),
    ("04", "인성/봉사", 20, "item-407", "교내 행사 질서/안전 도우미", 5, "누적합산형", "선택 (교사확인서)", "입학설명회, 축제, 체육대회 운영요원 성실 참여"),

    # 5. 창체/독서 (상한 15점)
    ("05", "창체/독서", 15, "item-501", "전공/인문 독서 활동 (10권 이상)", 8, "누적합산형", "필수 (독후감/서평)", "독서교육종합지원시스템 등록 또는 서평 보고서"),
    ("05", "창체/독서", 15, "item-502", "인문학 소양 아카데미 수료", 5, "고정점수형", "필수 (이수증)", "교내 인문학 특강 및 교양 프로그램 80% 수강"),
    ("05", "창체/독서", 15, "item-503", "창의적 체험활동 우수 산출물", 5, "누적합산형", "필수 (산출물 링크/파일)", "자율/진로활동 탐구보고서 발표 및 우수작 선정"),
    ("05", "창체/독서", 15, "item-504", "동아리 축제 부스 운영", 5, "고정점수형", "필수 (활동사진/보고서)", "교내 마이스터제 축제 전공/취미 동아리 부스 주도"),
    ("05", "창체/독서", 15, "item-505", "진로 탐색 및 기업 탐방 참여", 5, "누적합산형", "필수 (소감문)", "산업체 현장견학 및 IT 기업 멘토링 참여 보고서"),
    ("05", "창체/독서", 15, "item-506", "문화예술 / 체육활동 참가", 5, "누적합산형", "선택 (참가증빙)", "교내외 예술 동아리 발표 또는 스포츠 클럽 대회 참가"),
    ("05", "창체/독서", 15, "item-507", "자기주도학습 플래너 우수자", 4, "고정점수형", "선택 (플래너 검사)", "학기별 자기주도 학습 플래너 작성 성실 이수자")
]

row_start = 5
for item in criteria_33:
    r = ws5.max_row + 1
    ws5.append(list(item))
    ws5.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws5.cell(row=r, column=2).alignment = Alignment(horizontal="center", vertical="center")
    ws5.cell(row=r, column=2).font = data_font_bold
    ws5.cell(row=r, column=3).alignment = Alignment(horizontal="center", vertical="center")
    ws5.cell(row=r, column=4).font = data_font_code
    ws5.cell(row=r, column=4).alignment = Alignment(horizontal="center", vertical="center")
    ws5.cell(row=r, column=5).font = data_font_bold
    ws5.cell(row=r, column=6).alignment = Alignment(horizontal="center", vertical="center")
    ws5.cell(row=r, column=7).alignment = Alignment(horizontal="center", vertical="center")
    ws5.cell(row=r, column=8).alignment = Alignment(horizontal="center", vertical="center")
    ws5.cell(row=r, column=9).font = data_font
    ws5.cell(row=r, column=9).alignment = Alignment(wrap_text=True, vertical="center")

style_table(ws5, 1, 9, row_start, ws5.max_row)
auto_width(ws5, {"A": 10, "B": 14, "C": 14, "D": 14, "E": 28, "F": 12, "G": 18, "H": 18, "I": 45})

# Save Workbook
output_filename = "MEISTER_SYSTEM_FULL_SPEC.xlsx"
wb.save(output_filename)
print(f"Successfully generated {output_filename}")
